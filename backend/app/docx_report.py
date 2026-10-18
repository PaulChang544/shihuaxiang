"""通讯稿 + 活动照片 → Word 留档（python-docx）。"""

import datetime
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "微软雅黑"
MUTED = RGBColor(0x8A, 0x7D, 0x70)


def _safe_name(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "", text)[:40] or "活动"


def _apply_style_font(style) -> None:
    """统一样式字体（含中文 eastAsia），标题不再走英文 Light 字型。"""
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def render(out_path, images_root, act: dict, lead_name: str, report_text: str, photos: list) -> None:
    """act: activities 行（含 name/event_date/location/participants）；
    images_root: stored_path 的相对根（如 workspace 根，stored_path 形如 storage/uploads/...）；
    photos: uploads(category='photo') 行（stored_path/caption/filename）。"""
    doc = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        try:
            _apply_style_font(doc.styles[name])
        except KeyError:
            continue

    h = doc.add_heading(act["name"], level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta_bits = [
        act["event_date"] or "日期待定",
        act["location"] or "地点待定",
        f"负责人：{lead_name or '—'}",
    ]
    if str(act["participants"] or "").strip():
        meta_bits.append(f"参与：{act['participants']}")
    meta = doc.add_paragraph("　|　".join(meta_bits))
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.runs[0].font.color.rgb = MUTED
    meta.runs[0].font.size = Pt(9)

    for line in (report_text or "").splitlines():
        if line.strip():
            doc.add_paragraph(line)

    files = [
        p for p in photos
        if (p["stored_path"]) and p["stored_path"].strip()
    ]
    if files:
        doc.add_heading("活动留影", level=2)
        for p in files:
            path = images_root / p["stored_path"]
            if not path.exists():
                continue
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                pic.add_run().add_picture(str(path), width=Pt(380))
            except Exception:
                continue
            cap = (p["caption"] or p["filename"] or "").strip()
            if cap:
                cp = doc.add_paragraph(cap)
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.runs[0].font.color.rgb = MUTED
                cp.runs[0].font.size = Pt(9)

    doc.save(out_path)


def out_filename(act: dict) -> str:
    """文件名带活动 id：同名活动同日导出不会互相覆盖。"""
    return f"{_safe_name(act['name'])}-{datetime.date.today().isoformat()}-{act['id']}.docx"
